# -*- coding: utf-8 -*-
"""Rellena la plantilla Formato_Informe_LAB_3.docx con el informe del Laboratorio 3.

Uso:  python3 generar_informe_desde_plantilla.py PLANTILLA.docx ARBOL_INSTRUCCION.png EVIDENCIAS_DIR SALIDA.docx

Las derivaciones se calculan con herramientas/derivacion.py y las posiciones de error con
evidencias/validacion.json, de modo que el informe siempre coincide con la gramática y el validador.
El árbol con raíz <programa> se toma de ARBOL_INSTRUCCION.png con el sufijo _programa.png.
"""
import sys, json, re
from copy import deepcopy
from pathlib import Path
import docx
from docx.shared import Pt, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

TEMPLATE, TREE_PNG, EVID_DIR, OUT = map(Path, sys.argv[1:5])
TREE_PROG_PNG = TREE_PNG.with_name(TREE_PNG.stem + "_programa.png")
ROOT = EVID_DIR.resolve().parent
if not (ROOT / "herramientas").exists():   # disposición de la carpeta de Drive
    ROOT = ROOT / "gramática y validador"
sys.path.insert(0, str(ROOT / "herramientas"))
from derivacion import derivar_texto, formatear
def deriv(texto, desde, extension=False):
    return formatear(derivar_texto(texto, fragmento=True, extension=extension, desde=desde)).split("\n")
d = docx.Document(str(TEMPLATE))
body = d.element.body
E = list(body.iterchildren())          # referencias a los elementos originales de la plantilla
P = lambda i: Paragraph(E[i], d)
T = lambda i: Table(E[i], d)
FONT = "Times New Roman"

val = json.loads((EVID_DIR / "validacion.json").read_text(encoding="utf-8"))
casos = {c["id"]: c for c in val["casos"]}
def pos(cid):
    e = casos[cid]["error"]; return f"línea {e['linea']}, columna {e['columna']}"

# ------------------------------------------------------------------ utilidades
def style_run(r, size=None, bold=None, italic=None, mono=False):
    r.font.name = "Courier New" if mono else FONT
    rpr = r._r.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is not None:
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Courier New" if mono else FONT)
    if size: r.font.size = Pt(size)
    if bold is not None: r.bold = bold
    if italic is not None: r.italic = italic

def set_runs(par, runs, size=None, mono=False, keep_fmt=False):
    """runs: str o lista de (texto, {'b':bool,'i':bool,'m':bool})."""
    if isinstance(runs, str): runs = [(runs, {})]
    for r in list(par.runs): r._r.getparent().remove(r._r)
    for text, f in runs:
        r = par.add_run(text)
        style_run(r, size=size, bold=f.get("b"), italic=f.get("i"), mono=f.get("m", mono))
    return par

def new_par(after, runs, style="Normal", size=None, mono=False, space_after=None, align=None, left_indent=None, keep_next=False):
    p = d.add_paragraph(style=style)
    set_runs(p, runs, size=size, mono=mono)
    pf = p.paragraph_format
    if mono: pf.space_after = Pt(0); pf.space_before = Pt(0)
    if space_after is not None: pf.space_after = Pt(space_after)
    if left_indent is not None: pf.left_indent = Inches(left_indent)
    if align: p.alignment = align
    if keep_next: pf.keep_with_next = True
    after.addnext(p._p)
    return p._p

def chain(after, items, **kw):
    """Inserta varios párrafos seguidos; devuelve el último elemento."""
    cur = after
    for it in items:
        cur = new_par(cur, it, **kw)
    return cur

def code_block(after, lines, size=8.5):
    cur = after
    for k, ln in enumerate(lines):   # keep_next evita que un bloque de derivación se parta entre páginas
        cur = new_par(cur, ln, mono=True, size=size, left_indent=0.15, keep_next=(k < len(lines) - 1))
    return new_par(cur, "", space_after=2)

def remove(*idx):
    for i in idx:
        el = E[i]; el.getparent().remove(el)

def cell_set(cell, runs, size=9.5, bold=None, mono=False):
    if isinstance(runs, str): runs = [(runs, {})]
    paras = cell.paragraphs
    first = paras[0]
    for extra in paras[1:]: extra._p.getparent().remove(extra._p)
    for r in list(first.runs): r._r.getparent().remove(r._r)
    for text, f in runs:
        r = first.add_run(text)
        style_run(r, size=size, bold=f.get("b", bold), italic=f.get("i"), mono=f.get("m", mono))
    first.paragraph_format.space_after = Pt(1); first.paragraph_format.space_before = Pt(1)

def fit_rows(table, n_data):
    """Deja la tabla con exactamente n_data filas de datos (más el encabezado)."""
    rows = table.rows
    while len(table.rows) - 1 > n_data:
        tr = table.rows[-1]._tr; tr.getparent().remove(tr)
    while len(table.rows) - 1 < n_data:
        table._tbl.append(deepcopy(table.rows[-1]._tr))

def set_widths(table, inches):
    emus = [int(w * 914400) for w in inches]
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), emus): gc.set(qn("w:w"), str(int(w / 635)))
    for row in table.rows:
        for c, w in zip(row.cells, emus): c.width = Emu(w)

def fill(table, rows, widths=None, size=9.5, mono_cols=()):
    fit_rows(table, len(rows))
    for r, vals in zip(table.rows[1:], rows):
        for j, (c, v) in enumerate(zip(r.cells, vals)):
            cell_set(c, v, size=size, mono=(j in mono_cols))
    if widths: set_widths(table, widths)

M = {"m": True}  # marca de monoespaciado en una corrida

def espacio_antes(el, pts=5):
    Paragraph(el, d).paragraph_format.space_before = Pt(pts)
    return el


B = {"b": True}
NPRUEBAS = re.search(r"Ran (\d+) tests", (EVID_DIR / "pruebas.txt").read_text(encoding="utf-8")).group(1)
NCASOS = len(val["casos"])

# ------------------------------------------------------------------ encabezado
set_runs(P(1), "Gramáticas Libres de Contexto", size=16)
for run in P(1).runs: run.bold = True
t0 = T(3)
for r, (a, b) in zip(t0.rows[1:], [("Asignatura", "Teoría de la Computación – INFO1148"), ("Carrera", "Ingeniería Civil en Informática"),
                                  ("Docente", "Marcos Lévano"), ("Integrantes", "Juan Cárdenas y Bayron Flores (Jefe de Grupo)"),
                                  ("Fecha", "21 de octubre de 2026")]):
    cell_set(r.cells[0], a, size=10.5); cell_set(r.cells[1], b, size=10.5)

# ------------------------------------------------------------------ 1. Introducción
set_runs(P(5), "El objetivo del laboratorio es diseñar, analizar y validar una Gramática Libre de Contexto (GLC) que represente la "
  "estructura sintáctica de un lenguaje de programación simplificado. Se identifican los terminales, los no terminales y el símbolo "
  "inicial; se define formalmente G = (V, Σ, P, S); se construyen las producciones para programas con una o más instrucciones "
  "(asignaciones simples, asignaciones con una operación aritmética y salidas); y se comprueba la coherencia de la gramática mediante "
  "derivaciones por la izquierda, un árbol de derivación y la clasificación de cadenas válidas e inválidas. Los conceptos centrales son "
  "símbolo inicial, lenguaje generado L(G), derivación, árbol sintáctico y la separación entre análisis léxico (identificadores y números) "
  "y análisis sintáctico (orden de las categorías). Como extensión creativa se incorporan declaraciones de variables de tipo entero y real.")

# ------------------------------------------------------------------ 2.1 Pregunta 1
fill(T(9), [
  ("Palabras reservadas", [("Begin, Programa, End, salida", M)]),
  ("Operadores", [("+  -  *  /  %", M), (" (aritméticos) y ", {}), ("=", M), (" (asignación)", {})]),
  ("Símbolos especiales", [("#", M), (" (terminador de instrucción)", {})]),
  ("Identificadores", [("ID", M), (": categoría léxica de nombres como x, z2 y operativo_01_02. Patrón: L (L | D | _)*, con L = letra y D = dígito.", {})]),
  ("Números", [("NUM", M), (": categoría léxica de valores como 20, 0 y 1.345. Patrón: D+ (. D+)?, sin signo.", {})]),
], widths=[1.6, 5.3])

set_runs(P(12), [("V = { <programa>, <lista_instrucciones>, <instruccion>, <asignacion_simple>, <asignacion_aritmetica>, <expresion>, "
                  "<operando>, <operador>, <salida> }", M)], size=9.5)
set_runs(P(13), "Se usan dos no terminales para las asignaciones (<asignacion_simple> y <asignacion_aritmetica>) en lugar de uno solo, "
                "porque el enunciado pide diferenciarlas explícitamente (punto 7); ambos cumplen el rol de <asignacion> sugerido.")
set_runs(P(15), [("S = <programa>", M)])
set_runs(P(17), [("G = (V, Σ, P, S)", {"b": True})])
set_runs(P(18), [("V = { <programa>, <lista_instrucciones>, <instruccion>, <asignacion_simple>, <asignacion_aritmetica>, <expresion>, "
                  "<operando>, <operador>, <salida> }   (9 no terminales)", M)], size=9.5)
set_runs(P(19), [("Σ = { Begin, Programa, End, salida, ID, NUM, =, #, +, -, *, /, % }   (13 terminales)", M)], size=9.5)
set_runs(P(20), [("P = ", {}), ("Ver producciones desarrolladas en la Pregunta 2", {}), (" (9 reglas que, al separar sus alternativas, suman 17 producciones).", {})])
set_runs(P(21), [("S = <programa>", M)])
_an = [
  "No se enumeran los identificadores ni los números dentro de la gramática porque son conjuntos infinitos o no acotados: existen "
  "nombres nuevos (x, z2, temporal, operativo_01_02…) y números de longitud arbitraria (20, 0, 1.345, 1000.25…). Listar solo algunos "
  "valores (por ejemplo ID → x | y | z) haría que la gramática rechazara programas correctos, y listarlos todos es imposible. Por eso se "
  "modelan como categorías léxicas con un patrón general (ID y NUM), y cada token conserva su lexema como dato asociado: en la notación "
  "ID[z], el terminal es ID y z es su lexema. No se añaden reglas como ID → x ni NUM → 20.",
  "Es una GLC porque el lado izquierdo de cada producción es un único no terminal. Decisiones léxicas: las palabras reservadas distinguen "
  "mayúsculas y minúsculas (un nombre como salida_total sigue siendo ID); NUM admite enteros y reales sin signo (el 0 aparece en el "
  "programa de referencia, así que se interpreta «positivos» como «no negativos»); los espacios y los comentarios /* … */ los descarta el "
  "analizador léxico y no forman parte de la GLC. Los patrones de ID y NUM se expresan además como producciones a nivel de carácter en el Anexo A.1.",
]
set_runs(P(23), _an[0]); P(23).paragraph_format.space_after = Pt(4)
chain(E[23], _an[1:], space_after=4)
remove(24)

# ------------------------------------------------------------------ 2.2 Pregunta 2
prod = [
  ("<programa>", "<programa> → Begin Programa ID <lista_instrucciones> End"),
  ("<lista_instrucciones>", "<lista_instrucciones> → <instruccion> <lista_instrucciones> | <instruccion>"),
  ("<instruccion>", "<instruccion> → <asignacion_simple> | <asignacion_aritmetica> | <salida>"),
  ("<asignacion_simple>", "<asignacion_simple> → ID = <operando> #"),
  ("<asignacion_aritmetica>", "<asignacion_aritmetica> → ID = <expresion> #"),
  ("<expresion>", "<expresion> → <operando> <operador> <operando>"),
  ("<operando>", "<operando> → ID | NUM"),
  ("<operador>", "<operador> → + | - | * | / | %"),
  ("<salida>", "<salida> → salida ID #"),
]
fill(T(27), [(a, [(b, M)]) for a, b in prod], widths=[1.9, 5.0], size=9, mono_cols=(0,))
chain(E[27], [
  [("Diferenciación explícita de las tres instrucciones", {"b": True})],
], space_after=2, keep_next=True)
Paragraph(E[27].getnext(), d).paragraph_format.space_before = Pt(6)
cur = E[27].getnext()
cur = chain(cur, [
  [("Asignación simple: ", {"b": True}), ("ID = <operando> #", M), (". Copia un identificador o un número a un destino, sin operador (x = 20 #, temporal = x #).", {})],
  [("Asignación con expresión aritmética: ", {"b": True}), ("ID = <operando> <operador> <operando> #", M), (". Contiene exactamente un operador entre dos operandos (z = x + y #).", {})],
  [("Instrucción de salida: ", {"b": True}), ("salida ID #", M), (". Empieza con la palabra reservada salida y no lleva el operador = (salida x #).", {})],
  "La gramática no es ambigua: tras ID = <operando>, el siguiente terminal decide la regla (# cierra una asignación simple; un operador "
  "inicia una asignación aritmética) y salida solo puede iniciar la tercera clase. La recursión de <lista_instrucciones> exige al menos "
  "una instrucción y permite cualquier cantidad de ellas.",
], style="Normal", size=10, space_after=3)
for el in [E[27].getnext().getnext(), E[27].getnext().getnext().getnext(), E[27].getnext().getnext().getnext().getnext()]:
    Paragraph(el, d).style = d.styles["List Bullet"]

nt = [
  ("<programa>", "Símbolo inicial. Construye un programa completo: cabecera Begin Programa ID, cuerpo y cierre End."),
  ("<lista_instrucciones>", "Produce una o más instrucciones en orden (recursión por la derecha); nunca deriva ε."),
  ("<instruccion>", "Selecciona el tipo de instrucción: asignación simple, asignación aritmética o salida."),
  ("<asignacion_simple>", "Asigna un operando (ID o NUM) a un identificador destino y termina con #."),
  ("<asignacion_aritmetica>", "Asigna al destino el resultado sintáctico de una expresión de una operación binaria."),
  ("<expresion>", "Une un operando izquierdo, un operador y un operando derecho."),
  ("<operando>", "Selecciona un identificador (ID) o un número (NUM)."),
  ("<operador>", "Selecciona uno de los cinco operadores aritméticos: + - * / %."),
  ("<salida>", "Forma la instrucción salida ID #, que imprime el valor de un identificador."),
]
# copia de la tabla de 2 columnas para el Anexo A.2 (se toma antes de rellenarla)
TABLA_ANEXO = deepcopy(E[30])
fill(T(30), [([(a, M)], b) for a, b in nt], widths=[1.9, 5.0], size=9)
set_runs(P(33),
  "Sí. El programa de referencia tiene la cabecera Begin Programa operativo_01_02, veinte instrucciones y el cierre End. Tras descartar "
  "los espacios y el comentario /* operaciones y salidas */ (que elimina el analizador léxico; la GLC no lo genera), cada instrucción "
  "coincide con una producción: 6 asignaciones simples (x = 1.345 #, y = 20 #, temporal = 0 #, temporal = x #, x = y #, y = temporal #), "
  "7 asignaciones aritméticas (z y z2 con +; h y h2 con -; k con *; t con /; m con %) y 7 salidas (z, h, k, t, m, x, y). La lista se "
  "deriva con 20 aplicaciones de <lista_instrucciones>. El validador acepta el archivo completo ejemplos/programa_referencia.txt "
  "(91 tokens y 20 nodos <instruccion> en el árbol). Alcance: cada asignación aritmética tiene una sola operación binaria; no hay paréntesis, "
  "operaciones encadenadas ni signo unario (x = a + b * c # y x = -1 # no pertenecen a L(<instruccion>)).")
pert = new_par(E[33], [("Pertenencia a L(<programa>) y a L(<instruccion>). ", {"b": True}),
  ("Como S = <programa>, el lenguaje generado es L(G) = L(<programa>): solo programas completos (Begin Programa … End). Una instrucción "
   "suelta, como z = x + y #, pertenece a L(<instruccion>) pero no a L(<programa>). Para una instrucción w se cumple que w ∈ L(<instruccion>) "
   "si y solo si Begin Programa prueba w End ∈ L(<programa>). En las Preguntas 3 y 4 cada cadena se comprueba contra el lenguaje "
   "que se indica, y las derivaciones y los árboles se muestran tanto desde <instruccion> como desde el símbolo inicial <programa>.", {})],
  size=10, space_after=4)
remove(34)

# ------------------------------------------------------------------ 2.3 Pregunta 3
intro3 = new_par(E[35], "En cada paso se sustituye el no terminal situado más a la izquierda. De cada cadena se muestran dos derivaciones: "
  "desde <instruccion>, que prueba que pertenece a L(<instruccion>), y desde <programa> (símbolo inicial S), con la instrucción dentro del programa "
  "mínimo Begin Programa prueba … End. Las formas sentenciales conservan el terminal de categoría ID; el último paso es la sustitución léxica por los "
  "lexemas concretos (Anexo A.1), que elige el valor del token y no es una producción de P.", size=10, space_after=4)
REGLAS_ARIT = ("<instruccion> → <asignacion_aritmetica>; <asignacion_aritmetica> → ID = <expresion> #; <expresion> → <operando> <operador> <operando>; "
               "<operando> → ID; <operador> → +; <operando> → ID.")
REGLAS_PROG = "<programa> → Begin Programa ID <lista_instrucciones> End; <lista_instrucciones> → <instruccion>; "
REGLAS_SAL = "<instruccion> → <salida>; <salida> → salida ID #."

def bloque_derivacion(ancla, texto, reglas):
    cur = new_par(ancla, [("Desde ", {"b": True}), ("<instruccion>", {"b": True, "m": True}), (" (la cadena pertenece a L(<instruccion>)):", {"b": True})], size=9.5, space_after=2, keep_next=True)
    cur = code_block(cur, deriv(texto, "instruccion"))
    cur = new_par(cur, [("Reglas aplicadas, en orden: ", B), (reglas, {})], size=9.5, space_after=4)
    cur = new_par(cur, [("Desde ", {"b": True}), ("<programa>", {"b": True, "m": True}), (" (símbolo inicial S), dentro de Begin Programa prueba … End:", {"b": True})], size=9.5, space_after=2, keep_next=True)
    cur = code_block(cur, deriv(texto, "programa"))
    return new_par(cur, [("Reglas aplicadas: ", B), (REGLAS_PROG + reglas, {})], size=9.5, space_after=4)

set_runs(P(37), [("z = x + y #", M)], size=10)
bloque_derivacion(E[37], "z = x + y #", REGLAS_ARIT)
remove(38)
set_runs(P(40), [("salida z #", M)], size=10)
bloque_derivacion(E[40], "salida z #", REGLAS_SAL)
remove(41)

# árboles
set_runs(P(43), "Árbol de derivación de z = x + y # con raíz <instruccion> (Figura 1) y el mismo árbol dentro del programa mínimo, con raíz el símbolo inicial <programa> (Figura 2).", size=10)
pic1 = d.add_paragraph(); pic1.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic1.paragraph_format.keep_with_next = True
pic1.add_run().add_picture(str(TREE_PNG), width=Inches(3.9))
E[43].addnext(pic1._p)
cap1 = new_par(pic1._p, [("Figura 1. ", {"b": True}), ("Árbol de derivación de z = x + y # desde <instruccion>. No terminales en azul (nodos internos); terminales en verde (hojas). "
    "Generado con Graphviz a partir de la estructura que devuelve el validador. No se utilizó JFLAP.", {})], size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
pic2 = d.add_paragraph(); pic2.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic2.paragraph_format.keep_with_next = True
pic2.add_run().add_picture(str(TREE_PROG_PNG), width=Inches(5.9))
cap1.addnext(pic2._p)
cap2 = new_par(pic2._p, [("Figura 2. ", {"b": True}), ("Árbol de derivación desde <programa> (raíz = símbolo inicial S), con la instrucción dentro de Begin Programa prueba … End.", {})],
    size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
remove(44)

set_runs(P(46), [("En el árbol de la Figura 1 la raíz es ", {}), ("<instruccion>", M), (" y en el de la Figura 2 es ", {}), ("<programa>", M),
  (", el símbolo inicial S, que representa el programa completo. Los nodos internos son siempre símbolos no terminales: "
   "cada uno se expande en sus hijos aplicando exactamente una producción, y el conjunto de hijos, leído de izquierda a derecha, es el lado derecho de esa "
   "producción (por ejemplo, ", {}), ("<expresion>", M), (" tiene como hijos ", {}), ("<operando> <operador> <operando>", M), (" según ", {}),
  ("<expresion> → <operando> <operador> <operando>", M), (").", {})])
cur = chain(E[46], [
  [("Las hojas son los símbolos terminales y no tienen hijos. En la Figura 1, leídas de izquierda a derecha, forman la cadena derivada: ", {}),
   ("ID[z], =, ID[x], +, ID[y], #", M), (", es decir, z = x + y #. En la Figura 2 las hojas son ", {}),
   ("Begin, Programa, ID[prueba], ID[z], =, ID[x], +, ID[y], #, End", M),
   (", y forman el programa Begin Programa prueba z = x + y # End; la subcadena z = x + y # corresponde al subárbol de ", {}), ("<instruccion>", M),
   (". Cada árbol coincide paso a paso con la derivación por la izquierda correspondiente de la sección a): cada expansión de un no terminal "
    "es un nodo interno con sus hijos. Los corchetes ID[z] indican el lexema del token y no son símbolos adicionales.", {})],
], space_after=4)
remove(47)

fill(T(49), [
  ([("x = 20 #", M)], "Sí, ∈ L(<instruccion>)", "<instruccion> ⇒ <asignacion_simple> ⇒ ID = <operando> # ⇒ ID = NUM # ⇒ x = 20 #."),
  ([("z = x * y #", M)], "Sí, ∈ L(<instruccion>)", "<instruccion> ⇒ <asignacion_aritmetica> ⇒ ID = <expresion> # ⇒ ID = <operando> <operador> <operando> # ⇒ ID = ID <operador> <operando> # ⇒ ID = ID * <operando> # ⇒ ID = ID * ID # ⇒ z = x * y #. El operador * es una alternativa de <operador>."),
  ([("salida temporal #", M)], "Sí, ∈ L(<instruccion>)", "<instruccion> ⇒ <salida> ⇒ salida ID # ⇒ salida temporal #. temporal es un ID válido."),
], widths=[1.5, 1.4, 4.0], size=9)
nota = espacio_antes(new_par(E[49], "Las tres cadenas se comprueban contra L(<instruccion>). Por sí solas no pertenecen a L(<programa>), porque les faltan Begin Programa … End; "
  "dentro del programa mínimo Begin Programa prueba … End sí forman programas válidos de L(G).", size=9.5, space_after=6))

# ------------------------------------------------------------------ 2.4 Pregunta 4
intro4 = new_par(E[50], "Los casos 1 a 7 son instrucciones y se comprueban contra L(<instruccion>); el caso 8 es un programa completo y se comprueba contra "
  "L(<programa>). Ninguna de las instrucciones 1 a 3 pertenece a L(<programa>) por sí sola. Las posiciones de error (línea, columna) las informa "
  "validador.py (--fragmento para los casos 1 a 7).", size=10, space_after=4)
SI_I, NO_I = "Sí, ∈ L(<instruccion>)", "No, ∉ L(<instruccion>)"
rows4 = [
  ([("Caso 1", {"b": True}), ("\n", {}), ("x = 10 #", M)], SI_I, "Asignación simple: ID = NUM #. Todos los símbolos están en su lugar."),
  ([("Caso 2", {"b": True}), ("\n", {}), ("z = x + y #", M)], SI_I, "Asignación aritmética: ID = ID + ID #, con un operador entre dos operandos."),
  ([("Caso 3", {"b": True}), ("\n", {}), ("salida z #", M)], SI_I, "Salida: salida ID #."),
  ([("Caso 4", {"b": True}), ("\n", {}), ("x 10 #", M)], NO_I, f"Falta el operador de asignación = entre x y 10; tras el identificador se encontró el número 10 ({pos('4')})."),
  ([("Caso 5", {"b": True}), ("\n", {}), ("z = + x y #", M)], NO_I, f"Operador mal ubicado: después de = se requiere un operando (ID o NUM) y aparece + ({pos('5')}). La expresión es infija: <operando> <operador> <operando>."),
  ([("Caso 6", {"b": True}), ("\n", {}), ("salida #", M)], NO_I, f"Falta el identificador que debe seguir a salida; se encontró # ({pos('6')})."),
  ([("Caso 7", {"b": True}), ("\n", {}), ("x = 20", M)], NO_I, f"Falta el terminador #: la entrada termina donde se esperaba # o un operador ({pos('7')})."),
  ([("Caso 8", {"b": True}), ("\n", {}), ("Begin / Programa prueba / x = 5 # / salida x # / End", M)], "Sí, ∈ L(<programa>)",
   "Programa completo: Begin Programa ID <lista_instrucciones> End, con una asignación simple y una salida. Las barras del enunciado se leen como separadores de línea (ver nota bajo la tabla)."),
]
fill(T(51), rows4, widths=[1.7, 1.3, 3.9], size=9)
nota8 = espacio_antes(new_par(E[51], [("Nota sobre el caso 8. ", {"b": True}),
  ("El enunciado escribe el programa con barras / entre sus componentes. Se leen como separadores de línea, que es la única lectura bajo la cual la cadena es un programa; "
   "el programa queda así:", {})], size=9.5, space_after=2, keep_next=True))
cur8 = code_block(nota8, ["Begin", "Programa prueba", "x = 5 #", "salida x #", "End"], size=9)
cur8 = new_par(cur8, f"Si las barras fueran caracteres literales, el validador rechazaría la cadena en {pos('8_nota_barras_literales')}, porque aparece / donde la producción exige Programa. "
  "La barra no es separador del lenguaje: dentro de una expresión es división y /* inicia un comentario. Ambas lecturas constan en el Anexo A.2.", size=9.5, space_after=4)
set_runs(P(54), "Las cadenas incorrectas (casos 4 a 7) fallan por un único elemento sintáctico faltante o mal ubicado, y el análisis las detecta en el primer "
  "símbolo que ninguna producción puede continuar. Caso 4: falta el terminal =. Caso 5: el operador + está antes del primer operando (posición incorrecta; "
  "además x e y van después del operador, en notación prefija, y la gramática exige notación infija). Caso 6: falta el ID después de salida. Caso 7: falta el terminador #. "
  "Los casos válidos son: 1, 2 y 3, que pertenecen a L(<instruccion>) como asignación simple, asignación aritmética y salida, y el caso 8, que pertenece a L(<programa>) como programa completo. "
  "Todas las decisiones son sintácticas: no se evalúan valores ni se comprueba si una variable fue definida.")
remove(55)

# ------------------------------------------------------------------ 3. Extensión
set_runs(P(57), [("Construcción seleccionada: ", {}), ("Opción C – Declaración de variables", {"b": True})])
set_runs(P(59), [("entero x #", M), ("   y   ", {}), ("real temperatura #", M), (": una declaración es un tipo (entero o real), un identificador y el terminador #. "
                 "Es una instrucción más, que puede aparecer en cualquier posición de la lista de instrucciones.", {})])
remove(60)
fill(T(62), [
  ("Nuevos terminales", [("entero, real", M), (" (palabras reservadas, solo cuando la extensión está activa).", {})]),
  ("Nuevos no terminales", [("<declaracion>, <tipo>", M)]),
  ("Nuevas producciones", [("<declaracion> → <tipo> ID #", M), ("\n", {}), ("<tipo> → entero | real", M), ("\n", {}),
                          ("Modificada: <instruccion> → <asignacion_simple> | <asignacion_aritmetica> | <salida> | <declaracion>", M)]),
], widths=[1.9, 5.0], size=9)
set_runs(P(65), [("entero edad #", M), (" (también real temperatura #). Es válida porque tiene tipo, identificador y terminador.", {})])
casos_ext = casos["extension_invalida"]["error"]
set_runs(P(67), [("entero #", M), (f": inválida; falta el identificador después del tipo (línea {casos_ext['linea']}, columna {casos_ext['columna']}).", {})])
cur = code_block(E[68], deriv("entero edad #", "instruccion", extension=True))
cur = new_par(cur, "Resultado: entero edad # ∈ L(<instruccion>) y, dentro de Begin Programa prueba … End, el programa pertenece a L(G_ext). Las demás producciones y el símbolo inicial no cambian; "
   "G_ext = (V ∪ {<declaracion>, <tipo>}, Σ ∪ {entero, real}, P_ext, <programa>). La extensión define solo sintaxis: no comprueba "
   "declaraciones repetidas, declaración previa ni compatibilidad de tipos. El archivo ejemplos/programa_extendido.txt (declaraciones, "
   "asignaciones y salidas) es aceptado por el validador con --extension, y la derivación desde <programa> consta en evidencias/derivaciones.txt.", size=9.5, space_after=4)
remove(69)

# ------------------------------------------------------------------ 4. Conclusiones
set_runs(P(71), "La GLC propuesta, con S = <programa>, representa programas con una o más instrucciones y distingue asignaciones simples, asignaciones "
  "con una operación aritmética y salidas. Se comprobó que ID y NUM, definidos como categorías léxicas, permiten reutilizar las mismas producciones "
  "para cualquier nombre o valor, y que las veinte instrucciones del programa de referencia quedan cubiertas. Un aprendizaje importante fue "
  "que la pertenencia depende del símbolo desde el que se comprueba: las instrucciones sueltas pertenecen a L(<instruccion>) pero no a L(<programa>), "
  "y por eso cada resultado indica contra qué lenguaje se verificó.")
cur = new_par(E[71], "Las derivaciones por la izquierda y el árbol de derivación mostraron cómo cada paso aplica una producción y cómo las hojas reconstruyen "
  "la cadena. La validación de los ocho casos mostró que los errores se localizan en un elemento sintáctico concreto (falta de =, operador mal ubicado, "
  "ID faltante, # faltante). La extensión de declaraciones se incorporó formalmente con dos terminales, dos no terminales y una producción modificada, "
  "manteniendo la coherencia con el resto de la gramática. Queda fuera del alcance la semántica (tipos, variables declaradas y evaluación).", space_after=4)
remove(72)

# ------------------------------------------------------------------ 5. Consideraciones de entrega + Anexo
for i in range(74, 81): remove(i)
sec5 = E[73]
cur = chain(sec5, [
  "Informe en formato PDF (con copia editable en DOCX), con la numeración de secciones de la plantilla (1 a 5).",
  "Derivaciones, árboles y validaciones se generan con la misma gramática (gramatica.bnf) y el mismo validador.",
  "No se utilizó JFLAP; los árboles se generaron con Graphviz (.dot, .svg y .png en el repositorio).",
], style="List Bullet", size=10, space_after=1)
cur = new_par(cur, [("Anexo – Evidencias", {"b": True})], size=11, space_after=2, keep_next=True)
cur = chain(cur, [
  [("Repositorio GitHub: ", {"b": True}), ("https://github.com/der-kaffe/info1148-laboratorio-3", {})],
  [("Carpeta de Drive: ", {"b": True}), ("https://drive.google.com/drive/folders/1bNQcTqB4XKDPTkW4AXRuGCijnVxjX7TX", {})],
  [("Contenido: ", {"b": True}), ("gramatica.bnf, validador.py, ejemplos/ (programa de referencia, programa extendido y casos), evidencias/ "
    f"(validacion.json, derivaciones.txt, árboles en JSON y pruebas.txt con {NPRUEBAS} pruebas automatizadas) y arbol_sintactico/ (.dot, .svg, .png).", {})],
], size=10, space_after=2)

# A.1 gramática léxica a nivel de carácter
cur = new_par(cur, [("A.1 Gramática léxica a nivel de carácter", {"b": True})], size=10.5, space_after=2, keep_next=True)
cur = new_par(cur, "Los terminales ID y NUM se definen también con producciones libres de contexto sobre caracteres, con ... como abreviatura de los rangos. "
  "ID es <id> con la restricción léxica de no coincidir con una palabra reservada, y NUM es <num>. Con estas reglas la sustitución léxica de las derivaciones "
  "es una derivación más; por ejemplo, <id> ⇒ <letra> <resto_id> ⇒ z <resto_id> ⇒ z.", size=9.5, space_after=3)
cur = code_block(cur, [
  "<id>       → <letra> <resto_id>",
  "<resto_id> → ε | <letra> <resto_id> | <digito> <resto_id> | _ <resto_id>",
  "<num>      → <digitos> | <digitos> . <digitos>",
  "<digitos>  → <digito> | <digito> <digitos>",
  "<letra>    → A | B | ... | Z | a | b | ... | z | Ñ | ñ | Á | É | Í | Ó | Ú | Ü | á | é | í | ó | ú | ü",
  "<digito>   → 0 | 1 | 2 | ... | 9",
], size=7.5)

# A.2 resultado del validador
cur = new_par(cur, [("A.2 Resultado del validador en los casos del enunciado", {"b": True})], size=10.5, space_after=2, keep_next=True)
def etiqueta_caso(c):
    nombres = {"8": "Caso 8, una línea por componente", "8_nota_barras_literales": "Caso 8, con / literales",
               "extra_dos_operadores": "Dos operadores", "extra_signo_unario": "Signo unario",
               "extra_lista_vacia": "Lista vacía", "extra_dos_instrucciones": "Dos instrucciones",
               "P3_instruccion_sola_no_es_programa": "Instrucción sola como programa"}
    return nombres.get(c["id"], f"Caso {c['id']}")
def respuesta(c):
    if c["resultado"]:
        return f"Válido: pertenece a {c['lenguaje']}."
    e = c["error"]
    return f"Inválido para {c['lenguaje']} (línea {e['linea']}, columna {e['columna']}): {e['mensaje']}"
ids = ["1", "2", "3", "4", "5", "6", "7", "8", "8_nota_barras_literales", "extra_dos_operadores", "extra_signo_unario",
       "extra_lista_vacia", "extra_dos_instrucciones", "P3_instruccion_sola_no_es_programa"]
filas = []
for i in ids:
    c = casos[i]
    texto = c["texto"].replace("\n", " ⏎ ")
    filas.append(([(etiqueta_caso(c), {"b": True}), ("\n", {}), (texto, M)], respuesta(c)))
cur.addnext(TABLA_ANEXO)
tabla = Table(TABLA_ANEXO, d)
fill(tabla, filas, widths=[2.2, 4.7], size=8.5)
from docx.oxml import OxmlElement
for fila in tabla.rows:   # una fila nunca se parte entre dos páginas
    trPr = fila._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))
for c, t in zip(tabla.rows[0].cells, ("Entrada", "Respuesta del validador")):
    cell_set(c, t, size=9, bold=True)
cur = TABLA_ANEXO

# A.3 reproducción
cur = espacio_antes(new_par(cur, [("A.3 Cómo repetir las comprobaciones", {"b": True})], size=10.5, space_after=2, keep_next=True), 8)
cur = new_par(cur, f"El comando python3 -m unittest discover -s pruebas ejecutó {NPRUEBAS} pruebas con resultado OK (detalle en evidencias/pruebas.txt), y generar_evidencias.py "
  f"reproduce las {NCASOS} comprobaciones de cadenas, los dos programas completos, las derivaciones y los árboles.", size=9.5, space_after=3)
cur = code_block(cur, [
  "python3 validador.py ejemplos/programa_referencia.txt",
  'python3 validador.py --fragmento --cadena "z = x + y #"',
  "python3 validador.py ejemplos/programa_extendido.txt --extension",
  'python3 herramientas/derivacion.py --fragmento --desde instruccion --cadena "z = x + y #"',
  "python3 -m unittest discover -s pruebas -v",
  "python3 herramientas/generar_evidencias.py",
], size=9)
cur = new_par(cur, [("Programa de referencia validado:", {"b": True})], size=10, space_after=2, keep_next=True)
ref = (ROOT / "ejemplos" / "programa_referencia.txt").read_text(encoding="utf-8").splitlines()
cur = code_block(cur, [l for l in ref if l.strip()], size=8)

d.core_properties.title = "Laboratorio 3 – Diseño y Validación de Gramáticas"
d.core_properties.author = "Juan Cárdenas y Bayron Flores"
d.save(str(OUT))
print("ok", OUT)
